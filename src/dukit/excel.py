
import os
import openpyxl
import pandas as pd

from .typing import list_
from .utils import (
    log,
    )


def format(
        path,
        sheet=None,
        freeze_panes='B2',
        hide_sheets=None,
        hide_cols=None,
        col_width_max=70,
        col_width_padding=2,
        align_vertical='top',
        align_horizontal='left',
        wrap_text=True,
        verbosity=3,
        ):
    """
        apply formatting to an Excel file in place.

        col widths are adjusted to cell content,
        text is aligned and wrapped, and requested
        sheets or cols can be hidden.

        Parameters
        ----------
        path : str or os.PathLike
            file to format.
        sheet : str or list[str], optional
            sheets to format. defaults to all sheets.
        freeze_panes : str, optional
            excel cell reference for frozen panes.
        hide_sheets : str or list[str], optional
            sheets to hide.
        hide_cols : str or list[str], optional
            col names to hide.
        col_width_max : int, default 70
            maximum calculated col width.
        col_width_padding : int, default 2
            extra width added to calculated widths.
        align_vertical : str, default 'top'
            vertical cell alignment.
        align_horizontal : str, default 'left'
            horizontal cell alignment.
        wrap_text : bool, default True
            whether cell contents should wrap.
        verbosity : int, default 3
            logging verbosity level.

        Returns
        -------
        None
            the workbook is saved back to ``path``.
    """

    wb = openpyxl.load_workbook(path)

    if sheet:
        sheetnames = list_(sheet)
    else:
        sheetnames = [sheetname for sheetname in wb.sheetnames]
    sheets_hide = list_(hide_sheets)
    cols_hide = list_(hide_cols)


    for sheetname in sheetnames:

        data = pd.read_excel(path, sheet_name=sheetname)
        sheet = wb[sheetname]

        if sheetname in sheets_hide:
            sheet.sheet_state = 'hidden'

        if freeze_panes:
            sheet.freeze_panes = freeze_panes

        #adjust column widths, cell alignment and text wrapping
        for col in sheet.columns:
            colname = col[0].value
            col_letter = col[0].column_letter  #type:ignore

            if colname is None:
                text = (
                    'warning: skipping column with'
                    f' no header in sheet "{sheet.title}"'
                    )
                log(text, 'dk.format_excel()', verbosity)
                continue

            if colname in cols_hide:
                sheet.column_dimensions[col_letter].hidden = True  #type:ignore


            #multiline cells are split by newline,
            #expanded into new rows,
            #and the maximum length is calculated
            #these changes are not applied to the actual data,
            #only the column width is adjusted
            max_length = (
                data[colname]  #type:ignore
                .apply(str)
                .apply(lambda x: x.split('\n'))
                .explode()
                .str
                .len()
                .max()
                )
            max_length = max(max_length, len(str(colname)))
            max_length = min(max_length, col_width_max) + col_width_padding

            sheet.column_dimensions[col_letter].width = max_length

            for cell in col:
                cell.alignment = openpyxl.styles.Alignment(  #type:ignore
                    vertical=align_vertical,
                    horizontal=align_horizontal,
                    wrap_text=wrap_text
                    )

    wb.save(path)
    wb.close()




def save(
        df: pd.DataFrame,
        path: str,
        sheet_name='df',
        index=False,
        format_excel=True,
        **kwargs,
        ):
    """
    write a df to an Excel file.

    Parameters
    ----------
    df : pandas.DataFrame
        df to write.
    path : str
        destination file path.
    sheet_name : str, default 'df'
        sheet name.
    index : bool, default False
        whether to write the df index.
    format_excel : bool, default True
        whether to format the workbook after writing.
    **kwargs
        additional keyword arguments passed to
        :meth:`pandas.DataFrame.to_excel`.

    Returns
    -------
    None
        the df is written to ``path``.
    """

    if os.path.exists(path):
        writer = pd.ExcelWriter(
            path,
            mode='a',
            if_sheet_exists='replace',
            )
    else:
        writer = pd.ExcelWriter(
            path,
            mode='w',
            )

    with writer:
        df.to_excel(
            writer,
            sheet_name=sheet_name,
            index=index,
            **kwargs,
            )

    if format_excel:
        format(path, sheet=sheet_name)
